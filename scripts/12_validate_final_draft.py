from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from docx import Document

from capstone.config import project_root
from capstone.utils import sha256_file, write_json


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if len(s.split()) >= 8]


def main() -> None:
    root = project_root()
    docx = root / "reports" / "final_draft" / "QM640_Capstone_Final_Report_Draft_Debodip_Gupta.docx"
    doc = Document(docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    words = re.findall(r"\b[\w'-]+\b", text)
    placeholders = sorted(set(re.findall(r"\[[^\]]+\]|<[^>]+>", text)))
    prohibited = [
        phrase for phrase in [
            "Write 2–4 paragraphs", "<Insert", "Public repository placeholder",
            "synthetic dry-run dataset using", "Preliminary Findings",
        ] if phrase.lower() in text.lower()
    ]
    normalized = [re.sub(r"\W+", " ", s.lower()).strip() for s in _sentences(text)]
    duplicate_sentences = [sentence for sentence, count in Counter(normalized).items() if count > 1]
    long_quotes = []
    for quoted in re.findall(r'[“\"]([^”\"]+)[”\"]', text):
        if len(quoted.split()) > 25:
            long_quotes.append(quoted)
    required_headings = {
        "Introduction", "Problem Statement", "Literature Survey", "Data Description",
        "Research Design and Methodology", "Results", "Discussion", "Limitations",
        "Recommendations and Application", "Conclusion", "References",
    }
    headings = {p.text for p in doc.paragraphs if p.style and p.style.style_id in {"Heading1", "Heading2"}}

    baseline = json.loads((root / "reports" / "submission_artifact_baseline.json").read_text(encoding="utf-8"))
    artifact_checks = []
    for item in baseline["files"]:
        path = root / item["path"]
        actual = sha256_file(path)
        artifact_checks.append({"path": item["path"], "expected": item["sha256"], "actual": actual, "unchanged": actual.upper() == item["sha256"].upper()})

    report = {
        "docx": str(docx.relative_to(root)),
        "sha256": sha256_file(docx),
        "bytes": docx.stat().st_size,
        "word_count": len(words),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "figures": len(doc.inline_shapes),
        "placeholders": placeholders,
        "prohibited_template_phrases": prohibited,
        "duplicate_long_sentences": duplicate_sentences,
        "quotes_over_25_words": long_quotes,
        "missing_required_headings": sorted(required_headings - headings),
        "prior_submission_artifacts": artifact_checks,
    }
    report["passed"] = (
        not placeholders
        and not prohibited
        and not duplicate_sentences
        and not long_quotes
        and not report["missing_required_headings"]
        and all(item["unchanged"] for item in artifact_checks)
        and report["tables"] >= 10
        and report["figures"] >= 6
    )
    output = root / "reports" / "final_draft" / "validation_and_originality_audit.json"
    write_json(report, output)
    print(json.dumps(report, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
