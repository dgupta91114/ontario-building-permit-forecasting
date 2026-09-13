"""Check the final document, PDF, evidence, and bounded local originality indicators.

This does not query a plagiarism database and cannot certify a similarity percentage.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document

from capstone.config import project_root
from capstone.final_submission import REFERENCES


def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


def main():
    root=project_root(); out=root/"reports/final"
    stem="QM640_Capstone_Final_Report_Debodip_Gupta"
    doc=Document(out/f"{stem}.docx"); pdf=pymupdf.open(out/f"{stem}.pdf")
    manuscript=(out/"manuscript.md").read_text(encoding="utf-8")
    pages=manuscript.split("<!-- PAGE -->")
    abstract=normalize(pages[1].split("# Abstract",1)[1])
    doc_text="\n".join(p.text for p in doc.paragraphs)
    pdf_text="\n".join(p.get_text() for p in pdf)
    body="\n".join(p for p in manuscript.splitlines() if p and not p.startswith(("#","[[","<!--")))
    sentences=[normalize(x).lower() for x in re.split(r"(?<=[.!?])\s+",body) if len(x.split())>=12]
    duplicates=[x for x,n in Counter(sentences).items() if n>1]
    template_prompts=[x for x in ["<Topic>","<Name of the Student>","Write 2–4 paragraphs","Insert architecture diagram","Final Report Draft","Minimum 10 Relevant Sources","Include Y and X"] if x in doc_text]
    # Table notes can repeat standard units; originality screening is limited to prose.
    long_quotes=[x for x in re.findall(r'[“"]([^”"]+)[”"]',body) if len(x.split())>25]
    required=["Abstract","Introduction","Background and Context","Problem Statement","Purpose of the Study",
              "Research Problems / Research Questions","Contributions and Expected Value","Literature Review",
              "Literature Review Approach","Summary of Key Literature","Materials and Method",
              "Research Hypotheses","Minimum Sample Size Calculations","Architecture Diagram/Workflow",
              "System Overview","Architecture Diagram","Workflow Components","Tools and Technologies",
              "Results","Model Performance","Implementation and User Benefit","Deployment Approach",
              "System Integration","User Interaction","Benefits to Users","Example Use Case",
              "Limitations and Further Improvements","Limitations","Impact of Limitations",
              "Future Improvements","Future Scope","Bibliography","References",
              "Appendix A: Data Dictionary","Appendix B: Reproducibility and Evidence Map"]
    headings={p.text for p in doc.paragraphs if p.style and p.style.style_id.startswith("Heading")}
    missing_headings=sorted(set(required)-headings)
    catalog=json.loads((out/"figure_table_catalog.json").read_text())
    caption_checks=[]
    for entry in catalog:
        # Captions must appear intact on their planned PDF page.
        p=pdf[entry["planned_page"]-1] if len(pdf)>=entry["planned_page"] else None
        okay=p is not None and normalize(entry["title"]) in normalize(p.get_text()) and entry["label"] in p.get_text()
        caption_checks.append({**entry,"on_expected_page":okay})
    missing_xrefs=[]
    for entry in catalog:
        label=entry["label"]
        if re.search(r"\d",label) and not label.startswith(("Table A","Table B")):
            num=label.split()[1]
            if label not in body and not re.search(rf"{label.split()[0]}s\s+\d+\s+and\s+{num}\b",body):
                # Table 1 is cited in the plural 'Tables 1 and 2'.
                if not re.search(rf"{label.split()[0]}s\s+{num}\s+and\s+\d+\b",body): missing_xrefs.append(label)
    blank=[]; bad_numbers=[]; overflow=[]; page_checks=[]
    for i,p in enumerate(pdf,1):
        txt=p.get_text(); words=p.get_text("words")
        if len(txt.split())<12: blank.append(i)
        if not any(w[4]==str(i) and w[1]<60 for w in words): bad_numbers.append(i)
        for w in words:
            if w[0]<69 or w[2]>p.rect.width-68 or (w[1]>=60 and w[3]>p.rect.height-68):
                overflow.append({"page":i,"word":w[4],"bbox":list(w[:4])})
        page_checks.append({"page":i,"words":len(txt.split()),"images":len(p.get_images()),"links":len(p.get_links())})
    urls={r[3] for r in REFERENCES}
    pdf_urls={l.get("uri") for p in pdf for l in p.get_links() if l.get("uri")}
    missing_urls=sorted(urls-pdf_urls)
    refs_author_keys=["Angelopoulos", "Bank of Canada", "Bergmeir", "Breiman", "Chen", "Cohen", "Diebold", "Goh", "Hyndman", "Lundberg", "Newey", "Shmueli", "Statistics Canada", "Tashman", "Zou"]
    uncited=[x for x in refs_author_keys if x not in body]
    format_checks={"a4_page":all(abs(p.rect.width-595.3)<1 and abs(p.rect.height-841.9)<1 for p in pdf),
                   "one_inch_margins":all(all(abs(getattr(s,k).inches-1)<.001 for k in ["top_margin","bottom_margin","left_margin","right_margin"]) for s in doc.sections),
                   "times_new_roman_12pt_body":all(r.font.name=="Times New Roman" and r.font.size.pt==12 for p in doc.paragraphs for r in p.runs if p.style and p.style.style_id=="Normal" and getattr(p.paragraph_format.line_spacing,"pt",None)==24),
                   "body_leading_points":24,"body_font_points":12}
    evidence=json.loads((out/"evidence/evidence_audit.json").read_text())
    frozen=[]
    for item in evidence["frozen_artifacts"]:
        actual=hashlib.sha256((root/item["path"]).read_bytes()).hexdigest()
        frozen.append({"path":item["path"],"unchanged":actual==item["sha256"]})
    audit={"checked_on":"2026-09-13","pdf":f"{stem}.pdf","pdf_sha256":hashlib.sha256((out/f"{stem}.pdf").read_bytes()).hexdigest(),
           "pdf_pages":len(pdf),"abstract_words":len(abstract.split()),"pdf_extracted_words":len(pdf_text.split()),
           "tables":len(doc.tables),"figures":len(doc.inline_shapes),"references":len(REFERENCES),
           "missing_required_headings":missing_headings,"template_prompts":template_prompts,
           "blank_pages":blank,"missing_page_numbers":bad_numbers,"out_of_bounds_words":overflow,
           "missing_caption_cross_references":missing_xrefs,"caption_checks":caption_checks,
           "missing_reference_links":missing_urls,"uncited_reference_authors":uncited,
           "format":format_checks,"frozen_evidence":frozen,"pages":page_checks,
           "originality":{"local_duplicate_prose_sentences":duplicates,"quotes_over_25_words":long_quotes,
                          "external_similarity_percentage":None,
                          "external_plagiarism_check":"Not completed. A free web checker was inspected, but no full-document scan or percentage was obtained. Local checks and web phrase spot-checks do not verify the 7% requirement."}}
    audit["document_and_evidence_checks_passed"]=(len(pdf)==40 and 300<=len(abstract.split())<=350 and len(doc.tables)==18
        and len(doc.inline_shapes)==6 and len(REFERENCES)>=10 and not any([missing_headings,template_prompts,blank,bad_numbers,overflow,missing_xrefs,missing_urls,uncited,duplicates,long_quotes])
        and all(x["on_expected_page"] for x in caption_checks) and all(x["unchanged"] for x in frozen))
    (out/"quality_audit.json").write_text(json.dumps(audit,indent=2),encoding="utf-8")
    (out/"report_text.txt").write_text("\n".join(line.rstrip() for line in pdf_text.splitlines())+"\n",encoding="utf-8")
    # A plain-text copy makes manual scans possible without PDF copy/paste problems.
    scan=out/"qa/plagiarism_check_batches"; scan.mkdir(parents=True,exist_ok=True)
    prose_paragraphs=[p.text for p in doc.paragraphs if p.text and not (p.style and p.style.style_id.startswith("Heading"))]
    chunks=[]; chunk=[]; count=0
    for p in prose_paragraphs:
        n=len(p.split())
        if count+n>900 and chunk: chunks.append("\n\n".join(chunk));chunk=[];count=0
        chunk.append(p);count+=n
    if chunk:chunks.append("\n\n".join(chunk))
    for i,chunk in enumerate(chunks,1):(scan/f"prose-{i:02}.txt").write_text("\n".join(line.rstrip() for line in chunk.splitlines())+"\n",encoding="utf-8")
    print(json.dumps({k:v for k,v in audit.items() if k not in ["pages","caption_checks","frozen_evidence"]},indent=2))
    if not audit["document_and_evidence_checks_passed"]: raise SystemExit(1)


if __name__=="__main__": main()
