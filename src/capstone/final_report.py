from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

from .config import project_root


TITLE = "Forecasting Ontario Residential Building-Permit Values Using Explainable Time-Series and Machine-Learning Models"


def _clear_document(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_font(run, size: float = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _configure_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        header._p.append(field)


def _hyperlink(paragraph, text: str, url: str) -> None:
    """Append a genuine external Word hyperlink with Times New Roman formatting."""
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Times New Roman")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "24")
    properties.extend([fonts, color, underline, size, size_cs])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _linked_paragraph(doc: Document, lead: str, url: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(lead)
    _set_font(r)
    _hyperlink(p, url, url)
    return p


def _enforce_times_new_roman(doc: Document) -> None:
    """Enforce the requested font in styles, paragraphs, tables, and hyperlinks."""
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = "Times New Roman"
        if style._element.rPr is not None:
            fonts = style._element.rPr.rFonts
            if fonts is not None:
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{attribute}"), "Times New Roman")
    for paragraph in list(doc.paragraphs) + [p for section in doc.sections for p in section.header.paragraphs]:
        for run in paragraph.runs:
            _set_font(run, size=run.font.size.pt if run.font.size else 12, bold=bool(run.bold), italic=bool(run.italic))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_font(run, size=run.font.size.pt if run.font.size else 9, bold=bool(run.bold), italic=bool(run.italic))


def _paragraph(doc: Document, text: str, *, bold_lead: str | None = None, center: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None if center else Inches(0.5)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        _set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        _set_font(r)
    else:
        r = p.add_run(text)
        _set_font(r)
    return p


def _heading(doc: Document, text: str, level: int = 1):
    style = next((s for s in doc.styles if s.style_id == f"Heading{level}"), None)
    p = doc.add_paragraph(text, style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        _set_font(run, bold=True)
    return p


def _caption(doc: Document, label: str, title: str, note: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1
    r = p.add_run(label)
    _set_font(r, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1
    r = p.add_run(title)
    _set_font(r, italic=True)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1
        r = p.add_run("Note. ")
        _set_font(r, italic=True)
        r = p.add_run(note)
        _set_font(r)


def _format_value(value, digits: int = 3) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        if abs(float(value)) >= 1000:
            return f"{float(value):,.1f}"
        return f"{float(value):.{digits}f}"
    return str(value)


def _table(doc: Document, frame: pd.DataFrame, columns: list[tuple[str, str]], digits: int = 3) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    grid_style = next((s for s in doc.styles if s.style_id == "TableGrid"), None)
    if grid_style is not None:
        table.style = grid_style
    table.autofit = True
    for i, (_, label) in enumerate(columns):
        table.rows[0].cells[i].text = label
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for i, (column, _) in enumerate(columns):
            cells[i].text = _format_value(row[column], digits)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    _set_font(run, size=9, bold=row_index == 0)


def _figure(doc: Document, path: Path, label: str, title: str, note: str, width: float = 6.5) -> None:
    _caption(doc, label, title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    _caption(doc, "", "", note)


def _add_references(doc: Document) -> None:
    references = [
        "Bank of Canada. (2026). Valet API: How-to guide. https://www.bankofcanada.ca/valet-api-how-to/",
        "Bergmeir, C., Hyndman, R. J., & Koo, B. (2018). A note on the validity of cross-validation for evaluating autoregressive time series prediction. Computational Statistics & Data Analysis, 120, 70–83. https://doi.org/10.1016/j.csda.2017.11.003",
        "Breiman, L. (2001). Random forests. Machine Learning, 45(1), 5–32. https://doi.org/10.1023/A:1010933404324",
        "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 785–794. https://doi.org/10.1145/2939672.2939785",
        "Diebold, F. X., & Mariano, R. S. (1995). Comparing predictive accuracy. Journal of Business & Economic Statistics, 13(3), 253–263. https://doi.org/10.1080/07350015.1995.10524599",
        "Friedman, J. H. (2001). Greedy function approximation: A gradient boosting machine. The Annals of Statistics, 29(5), 1189–1232. https://doi.org/10.1214/aos/1013203451",
        "Goh, B. H. (1996). Residential construction demand forecasting using economic indicators: A comparative study of artificial neural networks and multiple regression. Construction Management and Economics, 14(1), 25–34. https://doi.org/10.1080/01446199600000004",
        "Goh, B. H. (2000). Evaluating the performance of combining neural networks and genetic algorithms to forecast construction demand: The case of the Singapore residential sector. Construction Management and Economics, 18(2), 209–217. https://doi.org/10.1080/014461900370834",
        "Hyndman, R. J., & Koehler, A. B. (2006). Another look at measures of forecast accuracy. International Journal of Forecasting, 22(4), 679–688. https://doi.org/10.1016/j.ijforecast.2006.03.001",
        "Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems, 30, 4765–4774. https://proceedings.neurips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html",
        "Makridakis, S., Spiliotis, E., & Assimakopoulos, V. (2018). The M4 Competition: Results, findings, conclusion and way forward. International Journal of Forecasting, 34(4), 802–808. https://doi.org/10.1016/j.ijforecast.2018.06.001",
        "Newey, W. K., & West, K. D. (1987). A simple, positive semi-definite, heteroskedasticity and autocorrelation consistent covariance matrix. Econometrica, 55(3), 703–708. https://doi.org/10.2307/1913610",
        "Shmueli, G. (2010). To explain or to predict? Statistical Science, 25(3), 289–310. https://doi.org/10.1214/10-STS330",
        "Sing, M. C. P., Edwards, D. J., Liu, H. J., & Love, P. E. D. (2015). Forecasting private-sector construction works: VAR model using economic indicators. Journal of Construction Engineering and Management, 141(11), 04015037. https://doi.org/10.1061/(ASCE)CO.1943-7862.0001016",
        "Statistics Canada. (2026a). Building permits, by type of structure and type of work (Table 34-10-0292-01). https://doi.org/10.25318/3410029201-eng",
        "Statistics Canada. (2026b). Labour force characteristics, monthly, seasonally adjusted and trend-cycle (Table 14-10-0287-01). https://doi.org/10.25318/1410028701-eng",
        "Tashman, L. J. (2000). Out-of-sample tests of forecasting accuracy: An analysis and review. International Journal of Forecasting, 16(4), 437–450. https://doi.org/10.1016/S0169-2070(00)00065-0",
    ]
    for reference in references:
        match = re.search(r"(https://\S+)$", reference)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 2
        text = reference[:match.start()].rstrip() + " " if match else reference
        r = p.add_run(text)
        _set_font(r)
        if match:
            _hyperlink(p, match.group(1), match.group(1))


def generate_final_draft(template: Path | None = None, root: Path | None = None) -> Path:
    base = root or project_root()
    template_path = template or base / "reports" / "templates" / "QM_640_Interim_Report_Template_Working_Copy.docx"
    output = base / "reports" / "final_draft" / "QM640_Capstone_Final_Report_Draft_Debodip_Gupta.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(template_path)
    _clear_document(doc)
    _configure_document(doc)

    tables = base / "outputs" / "tables"
    figures = base / "outputs" / "figures"
    summary = pd.read_csv(tables / "rq1_summary_by_cma.csv")
    persistence = pd.read_csv(tables / "rq1_persistence.csv")
    coefficients = pd.read_csv(tables / "rq1_hac_regression.csv")
    metrics = pd.read_csv(tables / "rq3_model_metrics.csv")
    rq2 = pd.read_csv(tables / "rq2_incremental_value.csv").iloc[0]
    paired = pd.read_csv(tables / "rq3_paired_error_tests.csv")
    regional = pd.read_csv(tables / "rq4_regional_metrics.csv")
    coverage = pd.read_csv(tables / "rq4_interval_coverage_by_cma.csv")
    importance = pd.read_csv(tables / "rq4_permutation_importance.csv")
    selection = json.loads((tables / "model_selection.json").read_text(encoding="utf-8"))
    feature_audit = json.loads((base / "data" / "processed" / "feature_audit.json").read_text(encoding="utf-8"))
    cleaning = json.loads((base / "data" / "interim" / "cleaning_audit.json").read_text(encoding="utf-8"))
    sens_tables = base / "outputs" / "sensitivity" / "constant_dollars" / "outputs" / "tables"
    sens_metrics = pd.read_csv(sens_tables / "rq3_model_metrics.csv")
    sens_selection = json.loads((sens_tables / "model_selection.json").read_text(encoding="utf-8"))
    xgb_test = paired[paired["model_a"] == "xgboost"].iloc[0]

    # Title page
    for _ in range(3):
        doc.add_paragraph()
    p = _paragraph(doc, TITLE, center=True)
    for run in p.runs:
        _set_font(run, bold=True)
    _paragraph(doc, "Final Report Draft", center=True)
    _paragraph(doc, "Debodip Gupta", center=True)
    _paragraph(doc, "Walsh College", center=True)
    _paragraph(doc, "QM640: Data Analytics Capstone", center=True)
    _paragraph(doc, "Mentor: Dr. Sanhita Karmakar", center=True)
    _paragraph(doc, "Summer 2026 Term", center=True)
    _paragraph(doc, "August 16, 2026", center=True)
    doc.add_page_break()

    _heading(doc, "GitHub Repository and Reproducibility Access")
    _linked_paragraph(doc, "Project repository: ", "https://github.com/dgupta91114/ontario-building-permit-forecasting")
    _paragraph(doc, "The repository contains the frozen source manifest, configuration, cleaning and feature audits, processed five-CMA dataset, analysis modules, tests, model-comparison outputs, figures, and reproducibility report. Raw national tables are reproducibly downloadable from official endpoints and are omitted from version control because the extracted files exceed ten gigabytes. SHA-256 hashes in the source manifest identify the exact archives used for this report.")

    _heading(doc, "Introduction")
    _heading(doc, "Background and Context", 2)
    for text in [
        "Building permits are issued before most construction begins and therefore provide a useful, although imperfect, signal of near-term construction intentions. Permit values affect decisions made by builders, material suppliers, lenders, municipal planners, workforce planners, and public agencies. They do not measure completed housing or completed investment. Instead, they summarize authorized projects whose timing, revision, or cancellation may differ from eventual construction activity (Statistics Canada, 2026a).",
        "This study examines monthly residential building-permit values in Toronto, Hamilton, Kitchener-Cambridge-Waterloo, London, and Windsor. The five census metropolitan areas differ substantially in scale and volatility. Pooling them creates enough observations for structured model comparison, while explicit CMA indicators and regional error reporting prevent the Toronto series from silently defining performance for every region.",
        "Forecasting permit values is difficult because the series combine serial dependence, annual seasonality, economic conditions, policy changes, pandemic disruption, and isolated large projects. A model can fit those patterns historically yet fail on future months. For that reason, this study uses chronological validation, a strong seasonal-naive control, multiple error measures, regional reliability checks, and uncertainty intervals rather than relying on training fit or random cross-validation (Bergmeir et al., 2018; Tashman, 2000).",
        "The analysis compares a transparent elastic-net model with random forest and extreme gradient boosting (XGBoost). Tree ensembles can represent nonlinear relationships and interactions, but they are accepted only when they outperform the benchmark under a prespecified governance rule (Breiman, 2001; Chen & Guestrin, 2016). Explainability is addressed through permutation importance and SHAP, which describe model behavior without converting predictive associations into causal claims (Lundberg & Lee, 2017; Shmueli, 2010).",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Problem Statement")
    _paragraph(doc, "Official public permit tables are descriptive and do not directly supply reproducible next-month regional forecasts, controlled model comparisons, or explanations of forecast drivers. Informal extrapolation can be distorted by seasonal timing, regional scale, temporary disruptions, and a few high-value projects. The practical problem is to determine whether a reproducible forecasting system based solely on aggregated public data can improve upon a defensible seasonal benchmark while remaining reliable across five Ontario CMAs.")
    _heading(doc, "Purpose of the Study", 2)
    _paragraph(doc, "The purpose of this quantitative secondary-data study is to describe temporal patterns, forecast next-month residential building-permit values, test whether lagged unemployment and the overnight-rate target add practical predictive information, compare four model strategies, and explain the preferred model. The study evaluates prediction rather than causation.")

    _heading(doc, "Scope and Research Questions")
    _paragraph(doc, "The unit of analysis is one CMA-month. Official observations span January 2018 through May 2026, the newest common permit month available for the frozen extraction. Lag creation and the one-month target lead produce 440 analytic observations with target dates from February 2019 through May 2026. Values are expressed in thousands of Canadian dollars.")
    for rq in [
        "RQ1. What trend, seasonal, persistence, and volatility patterns characterize monthly residential building-permit values across the five CMAs?",
        "RQ2. Do the Bank of Canada overnight-rate target and Ontario unemployment rate improve forecast accuracy beyond permit history, seasonality, and CMA effects?",
        "RQ3. Which candidate model produces the most accurate and stable one-month-ahead forecasts across time and regions?",
        "RQ4. Which variables contribute most to the preferred model, and does forecast performance differ materially across CMAs?",
    ]:
        _paragraph(doc, rq, bold_lead=rq[:4])

    _heading(doc, "Literature Survey")
    _heading(doc, "Literature Review Approach", 2)
    _paragraph(doc, "The literature review prioritized peer-reviewed forecasting, construction-demand, validation, accuracy-metric, ensemble-learning, and explainability research, supplemented by official source documentation. Sources were retained when they directly supported a research question, model family, evaluation choice, or interpretation safeguard. Descriptive commentary without transferable methods was excluded.")
    for text in [
        "Construction forecasting studies show why external economic variables are plausible but not guaranteed to help. Goh (1996, 2000) found value in nonlinear methods for construction demand, while Sing et al. (2015) used economic indicators in a vector autoregression framework. These studies motivate testing both linear and nonlinear candidates. They do not justify presuming that macroeconomic inputs will improve a particular Ontario permit forecast, so RQ2 is framed as an incremental comparison.",
        "Forecast evaluation research strongly supports future-like testing. Tashman (2000) emphasizes out-of-sample evaluation, and Bergmeir et al. (2018) describe conditions under which cross-validation can be appropriate for autoregressive data. The present design consequently keeps time ordered, assigns all CMA observations from a month to the same fold, tunes only on pre-holdout months, and reserves the latest 12 target months for final evaluation.",
        "No single metric fully describes forecast quality. MAE preserves the units decision makers understand, RMSE places more weight on large errors, and MASE scales errors relative to a seasonal benchmark (Hyndman & Koehler, 2006). sMAPE and WAPE provide percentage-oriented summaries. Paired tests based on common target observations complement, rather than replace, the practical thresholds used for selection (Diebold & Mariano, 1995).",
        "Random forests average decorrelated trees, while gradient boosting builds a sequence of corrective trees (Breiman, 2001; Friedman, 2001). XGBoost regularizes that boosting process and offers controlled depth, subsampling, and learning-rate parameters (Chen & Guestrin, 2016). The present grids are deliberately small because the dataset contains only 440 rows and extensive searching could overfit validation folds.",
        "Permutation importance quantifies the loss in holdout accuracy when a feature is disrupted. SHAP distributes a model prediction among transformed features using an additive attribution framework (Lundberg & Lee, 2017). Both describe how the fitted model uses available information; neither demonstrates that changing a feature would cause permit values to change (Shmueli, 2010).",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Data Description")
    _heading(doc, "Sources and Frozen Extraction", 2)
    _paragraph(doc, "Statistics Canada Table 34-10-0292-01 supplies permit values by month, geography, building type, work type, and seasonal-adjustment/value basis. Table 14-10-0287-01 supplies the seasonally adjusted Ontario unemployment-rate estimate. Bank of Canada series V39079 supplies daily overnight-rate targets, summarized as the last observed target in each calendar month (Bank of Canada, 2026; Statistics Canada, 2026a, 2026b).")
    sources = pd.DataFrame([
        ["Statistics Canada 34-10-0292-01", "Residential permit value", "Monthly", "2018-01 to 2026-05"],
        ["Statistics Canada 14-10-0287-01", "Ontario unemployment rate", "Monthly", "2018-01 to 2026-05"],
        ["Bank of Canada V39079", "Overnight-rate target", "Daily to monthly", "2018-01 to 2026-05"],
    ], columns=["source", "measure", "frequency", "coverage"])
    _caption(doc, "Table 1", "Official Data Sources")
    _table(doc, sources, [("source", "Source"), ("measure", "Measure"), ("frequency", "Frequency"), ("coverage", "Coverage")])
    _caption(doc, "", "", "The raw archive manifest records official URLs, retrieval timestamps, and SHA-256 hashes.")

    _heading(doc, "Data Preparation and Quality Controls", 2)
    _paragraph(doc, f"The permit cleaner streamed {cleaning['building_permits']['source_rows']:,} national rows and retained 505 rows: 101 months for each selected CMA. Filters required Value of permits, Total residential, Types of work total, seasonally adjusted current values, dollar units, and a thousands scalar. The unemployment cleaner scanned {cleaning['unemployment']['source_rows']:,} rows and required Ontario, Unemployment rate, Total - Gender, age 15 years and over, seasonally adjusted data, and Statistics = Estimate. All three selected source series contained 101 months and no missing selected values.")
    _paragraph(doc, "One London permit value for April 2025 carried Statistics Canada status E, meaning use with caution. The observation was retained because it is an official estimate, is not missing, and no evidence identified it as erroneous. Large values were also retained: major projects are substantively meaningful events rather than automatic outliers. Their influence is addressed through robust error reporting, regional results, holdout plots, and the constant-dollar sensitivity analysis.")
    _paragraph(doc, f"Feature construction removed {feature_audit['rows_dropped']} rows because a 12-month lag history, lagged macroeconomic values, and the future target were required. The resulting {feature_audit['rows_after_complete_case_filter']} complete rows contain 88 targets per CMA. Targets are built within CMA, and every forecast origin precedes its target month.")

    dictionary = pd.DataFrame([
        ["permit_value_t_plus_1", "Next-month residential permit value", "Numeric, $000s"],
        ["permit_value_t", "Permit value observed at forecast origin", "Numeric, $000s"],
        ["lag_1, lag_3, lag_6, lag_12", "Within-CMA permit lags", "Numeric, $000s"],
        ["rolling_mean_3", "Three-month mean through origin", "Numeric, $000s"],
        ["rolling_sd_3", "Three-month standard deviation through origin", "Numeric, $000s"],
        ["seasonal_naive_pred", "Same target month one year earlier", "Numeric, $000s"],
        ["month_sin, month_cos", "Cyclical target-month encoding", "Numeric"],
        ["overnight_rate_lag_1", "Prior-month overnight-rate target", "Percent"],
        ["unemployment_lag_1", "Prior-month Ontario unemployment rate", "Percent"],
        ["cma", "Census metropolitan area", "Categorical"],
    ], columns=["variable", "definition", "type"])
    _caption(doc, "Table 2", "Analytic Data Dictionary")
    _table(doc, dictionary, [("variable", "Variable"), ("definition", "Definition"), ("type", "Type")])

    _heading(doc, "Research Design and Methodology")
    _heading(doc, "Forecasting Design", 2)
    for text in [
        "Each row represents information available at origin month t and predicts permit value at t + 1. Permit lags and rolling measures are calculated independently within each CMA. The seasonal-naive forecast for t + 1 uses the value from the same calendar month one year earlier, which is located at t - 11 when rows are indexed by origin month. Lagged macroeconomic variables are used to reduce publication-timing risk.",
        "The latest 12 target months form the untouched temporal holdout. Earlier months are used in expanding-window validation with 36 initial training months, 12-month validation blocks, and 12-month steps. CMA rows from the same target month never cross split boundaries. Model preprocessing is fitted within each fold so scaling and category handling cannot learn from future months.",
        "Candidate models are seasonal naive, elastic net, random forest, and XGBoost. Elastic net provides history-only and macro-augmented variants for RQ2. Random forest and XGBoost use the complete feature set. Predictions are constrained to be nonnegative. Hyperparameter grids are prespecified and bounded, and the fixed random seed is 640.",
        "The selection rule requires pooled MASE below 1.0, MAE below seasonal naive, no material degradation greater than 10% in three or more CMAs, and successful reproducibility checks. If no fitted candidate qualifies, seasonal naive is retained. This rule separates governance from retrospective preference for the lowest displayed metric.",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Evaluation and Statistical Methods", 2)
    _paragraph(doc, "MAE and RMSE summarize absolute error in thousands of dollars. MASE divides error by a within-CMA seasonal scale; values below one indicate improvement over the seasonal scaling reference. sMAPE and WAPE provide percentage views. RQ2 uses a paired t test, Wilcoxon signed-rank test, and a dependence-aware Diebold-Mariano-style statistic, but practical importance still requires at least 5% MAE improvement. RQ4 uses regional metrics and a one-way ANOVA of absolute holdout errors. Ninety-percent split-conformal intervals are calibrated only with pre-holdout observations.")

    _heading(doc, "Results")
    _heading(doc, "RQ1: Temporal and Regional Structure", 2)
    _caption(doc, "Table 3", "Current-Dollar Permit Summary by CMA")
    _table(doc, summary, [("cma", "CMA"), ("n", "n"), ("mean", "Mean"), ("median", "Median"), ("std", "SD"), ("coefficient_of_variation", "CV")])
    _paragraph(doc, "Toronto had the largest scale, with a mean monthly value of $1.230 billion, while Windsor had the smallest mean at $58.5 million. Relative volatility was highest in Hamilton (CV = 0.634) and lowest in Toronto (CV = 0.241). These scale differences support pooled models with explicit CMA handling and make percentage and scaled metrics essential alongside pooled dollar error.")
    _caption(doc, "Table 4", "Permit Persistence by CMA")
    _table(doc, persistence, [("cma", "CMA"), ("acf_lag_1", "Lag-1 ACF"), ("acf_lag_12", "Lag-12 ACF")])
    _paragraph(doc, "Lag-1 autocorrelation ranged from 0.025 in Kitchener-Cambridge-Waterloo to 0.556 in Windsor. Lag-12 autocorrelation was weak in Hamilton and negative in the other four CMAs. Persistence therefore differed materially by region; a uniform assumption of strong annual autocorrelation would not describe the observed series.")
    _figure(doc, figures / "rq1_trend_by_cma.png", "Figure 1", "Monthly Residential Permit Values by CMA", "Official seasonally adjusted current-dollar values; vertical scale is thousands of Canadian dollars.")
    _figure(doc, figures / "rq1_seasonal_profile.png", "Figure 2", "Median Seasonal Profile by CMA", "Points are median permit values for each calendar month across the official analysis period.")
    _paragraph(doc, "The pooled HAC regression explained 92.3% of variation, largely reflecting CMA scale and the current permit level. Current permit value was positively associated with next-month value (b = 0.315, p < .001), and the time index was positive (b = 575.3, p = .007). The sine and cosine seasonal terms were not significant in this pooled specification. These coefficients are descriptive associations and should not be interpreted causally.")

    _heading(doc, "RQ2: Incremental Value of Macroeconomic Predictors", 2)
    rq2_frame = pd.DataFrame([rq2])
    _caption(doc, "Table 5", "History-Only Versus Macro-Augmented Elastic Net")
    _table(doc, rq2_frame, [("n", "n"), ("mae_a", "Macro MAE"), ("mae_b", "History MAE"), ("improvement_pct_a_vs_b", "Improvement %"), ("paired_t_p_value", "Paired p"), ("dm_style_p_value", "DM-style p"), ("meets_practical_threshold", "Meets 5%")])
    _paragraph(doc, f"The macro-augmented elastic net reduced MAE by only {rq2['improvement_pct_a_vs_b']:.2f}%. The paired t test (p = {rq2['paired_t_p_value']:.3f}), Wilcoxon test (p = {rq2['wilcoxon_p_value']:.3f}), and DM-style test (p = {rq2['dm_style_p_value']:.3f}) did not indicate a reliable improvement. Because the gain also fell below the prespecified 5% threshold, RQ2 is answered negatively: these two lagged macroeconomic variables did not add practically meaningful forecast accuracy beyond permit history, seasonality, trend, and CMA effects.")

    _heading(doc, "RQ3: Model Accuracy and Stability", 2)
    _caption(doc, "Table 6", "Official 12-Month Holdout Model Comparison")
    _table(doc, metrics, [("model", "Model"), ("n", "n"), ("mae", "MAE"), ("rmse", "RMSE"), ("mase", "MASE"), ("smape_pct", "sMAPE %"), ("wape_pct", "WAPE %")])
    _paragraph(doc, f"XGBoost produced the lowest MAE ({selection['pooled_mae']:,.1f}) and MASE ({selection['pooled_mase']:.3f}). Its MAE was 53.8% below seasonal naive. The paired comparison was significant by the t test (p = {xgb_test['paired_t_p_value']:.4f}), Wilcoxon test (p = {xgb_test['wilcoxon_p_value']:.4f}), and DM-style test (p < .001). XGBoost was not materially worse than seasonal naive in any CMA and therefore satisfied the prespecified rule.")
    _paragraph(doc, "Random forest was close to XGBoost in MASE but had higher dollar error. Both elastic-net variants outperformed seasonal naive but were clearly weaker than the tree ensembles. The outcome supports a nonlinear relationship among recent permit levels, region, and lagged history, while the bounded tree depth and holdout evaluation limit—but do not eliminate—overfitting risk.")
    _figure(doc, figures / "selected_model_actual_vs_predicted.png", "Figure 3", "XGBoost Holdout Forecasts and Actual Values", "Each panel contains the same 12 target months. Differences in vertical scale reflect actual regional permit-value differences.")
    _figure(doc, figures / "selected_model_absolute_error.png", "Figure 4", "XGBoost Absolute Holdout Error Over Time", "Absolute errors are reported in thousands of Canadian dollars.")

    _heading(doc, "RQ4: Explainability, Regional Reliability, and Uncertainty", 2)
    selected_regional = regional[regional["model"] == selection["selected_model"]]
    _caption(doc, "Table 7", "XGBoost Holdout Performance by CMA")
    _table(doc, selected_regional, [("cma", "CMA"), ("n", "n"), ("mae", "MAE"), ("rmse", "RMSE"), ("mase", "MASE"), ("smape_pct", "sMAPE %"), ("wape_pct", "WAPE %")])
    _paragraph(doc, "XGBoost achieved MASE below one in every CMA. Toronto had the largest absolute MAE because its permit values were roughly an order of magnitude larger, but it also had a low regional MASE of 0.295. Kitchener-Cambridge-Waterloo had the highest regional MASE at 0.924. Absolute errors differed significantly across regions (ANOVA p < .001), confirming that pooled performance should not be interpreted as uniform reliability.")
    _caption(doc, "Table 8", "XGBoost 90% Interval Coverage by CMA")
    _table(doc, coverage, [("cma", "CMA"), ("n", "n"), ("coverage", "Coverage"), ("mean_interval_width", "Mean Width")])
    _paragraph(doc, f"Overall empirical coverage was {selection['holdout_interval_coverage']:.1%}, exceeding the nominal 90% level, but the conformal radius was ${selection['conformal_radius']:,.0f} thousand. Coverage was therefore conservative and intervals were wide, especially for Toronto. The intervals communicate genuine operational uncertainty; point forecasts should not be treated as precise commitments.")
    _caption(doc, "Table 9", "Top Permutation-Importance Features")
    _table(doc, importance.head(8), [("feature", "Feature"), ("mae_increase_mean", "Mean MAE Increase"), ("mae_increase_sd", "SD")])
    _paragraph(doc, "The current permit value was the dominant permutation feature, followed by the three-month rolling mean and permit lags at six, three, and one months. CMA and the 12-month lag also contributed. The lagged overnight rate had zero measured importance and lagged unemployment was slightly negative, consistent with RQ2. SHAP produced the same broad conclusion that recent permit history and regional scale dominate model behavior.")
    _figure(doc, figures / "rq4_permutation_importance.png", "Figure 5", "XGBoost Permutation Importance", "Importance is the increase in holdout MAE after permuting one original feature; negative values indicate no reliable benefit.")
    _figure(doc, figures / "rq4_shap_beeswarm.png", "Figure 6", "XGBoost SHAP Distribution", "SHAP values are on the fitted log1p target scale and describe model attribution rather than causal effects.")

    _heading(doc, "Constant-Dollar Sensitivity Analysis")
    _caption(doc, "Table 10", "Constant-Dollar Holdout Model Comparison")
    _table(doc, sens_metrics, [("model", "Model"), ("mae", "MAE"), ("rmse", "RMSE"), ("mase", "MASE"), ("smape_pct", "sMAPE %")])
    _paragraph(doc, f"The constant-dollar analysis retained all 440 observations and again selected XGBoost. Its MAE was {sens_selection['pooled_mae']:,.1f}, MASE was {sens_selection['pooled_mase']:.3f}, and interval coverage was {sens_selection['holdout_interval_coverage']:.1%}. The macro-augmented elastic net performed worse than its history-only counterpart in this sensitivity analysis. The agreement in selected model and feature narrative indicates that the main conclusion is not an artifact of expressing permits in current rather than constant dollars.")

    _heading(doc, "Discussion")
    for text in [
        "The evidence answers the four research questions with a coherent pattern. Permit series differ greatly in scale and regional persistence. Lagged macroeconomic variables did not provide the expected practical gain. Nonlinear tree ensembles substantially improved out-of-sample error, and XGBoost met the prespecified governance criteria. Its behavior was driven primarily by current and recent permit history rather than the macroeconomic features.",
        "The result does not mean interest rates and labour-market conditions are irrelevant to construction. Monthly provincial unemployment and a single policy-rate series may be too aggregated, delayed, or indirectly related to local permit timing to improve a short-horizon model after recent permit history is known. Alternative macroeconomic data, longer horizons, local labour measures, mortgage rates, or construction-cost indexes could be evaluated in future work, but introducing them after observing the holdout would change the prespecified study.",
        "Regional reliability requires nuance. XGBoost improved upon seasonal scaling in every CMA, yet Kitchener-Cambridge-Waterloo and London were harder to predict on a scaled basis than Toronto and Windsor. Toronto dominated pooled dollar error because of its scale. Stakeholders should therefore use region-specific performance and intervals, not pooled MAE alone, when deciding whether a forecast is actionable.",
        "The conservative intervals are an important counterweight to the strong point-error improvement. They show that individual monthly permit values remain uncertain, particularly when large projects enter the series. The model is better suited to planning ranges, scenario discussion, and early-warning support than to deterministic budgeting.",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Limitations")
    for text in [
        "The analysis covers only five Ontario CMAs and a period beginning in 2018. Eight years of monthly data are sufficient for the prespecified comparisons but include the pandemic and relatively few independent annual cycles. Findings may not generalize to other regions, permit categories, or forecast horizons.",
        "Building permits represent authorized construction intentions, not starts, completions, housing supply, or realized investment. Revisions and large projects can change monthly values. Seasonal adjustment is performed by Statistics Canada, and one retained London observation carries an estimated-value caution flag.",
        "The holdout contains 60 CMA-month observations, only 12 per region. Statistical tests on paired errors have limited power and may be affected by temporal and cross-regional dependence. The ANOVA is descriptive evidence of regional error heterogeneity rather than a causal regional comparison.",
        "Model selection and reported point performance use the same final holdout governance exercise. The holdout was not used for hyperparameter fitting, but future prospective forecasts would provide stronger external validation. SHAP and permutation importance explain the fitted model and are not causal evidence.",
        "Raw official tables are extremely large. Reproducibility depends on scripted retrieval, chunked filtering, hashes, and source availability. Statistics Canada may revise historical seasonally adjusted values, so future downloads may differ from the frozen archive hashes used here.",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Recommendations and Application")
    for text in [
        "Use XGBoost as the preferred one-month-ahead planning model for the studied configuration, while retaining seasonal naive as a continuously monitored control. Refit only after the latest official month is available and evaluate every new forecast against the benchmark.",
        "Present forecasts by CMA with the conformal interval and regional historical error. Do not communicate a point value without its uncertainty range. Escalate unusually large forecasts or residuals for contextual review rather than automatically deleting them.",
        "Do not retain the unemployment and overnight-rate variables on the claim that they improve accuracy. They may remain in monitored research variants, but the current evidence does not satisfy the practical threshold. Simpler history-only inputs should be considered when operational cost or data latency matters.",
        "Maintain the frozen source manifest, configuration, package versions, tests, and generated evidence for every report release. Changes to source labels, model grids, horizon, or decision rules should be versioned and evaluated prospectively rather than adjusted to improve a completed holdout result.",
        "The intended users are analysts and planners who understand that permits are an early indicator. The system should support scenario planning and resource discussions, not replace professional judgment, municipal knowledge, or project-level review.",
    ]:
        _paragraph(doc, text)

    _heading(doc, "Conclusion")
    _paragraph(doc, "A reproducible official-data pipeline was used to forecast one-month-ahead residential building-permit values for five Ontario CMAs. XGBoost materially outperformed the seasonal-naive benchmark and met all prespecified selection conditions, while macroeconomic augmentation did not provide meaningful incremental accuracy. Recent permit history and regional scale dominated model behavior. Constant-dollar results supported the same model conclusion. The preferred model can assist short-term planning when forecasts are reported with regional performance and wide uncertainty intervals, but it should not be interpreted as a causal model or a precise forecast of completed construction.")

    doc.add_page_break()
    _heading(doc, "References")
    _add_references(doc)

    doc.add_page_break()
    _heading(doc, "Appendix A")
    _heading(doc, "Reproducibility and Evidence Map", 2)
    evidence = pd.DataFrame([
        ["Frozen source manifest", "data/raw/download_manifest.json"],
        ["Cleaning audit", "data/interim/cleaning_audit.json"],
        ["Feature audit", "data/processed/feature_audit.json"],
        ["Official analytic dataset", "data/processed/modeling_dataset.csv"],
        ["RQ1 outputs", "outputs/tables/rq1_* and outputs/figures/rq1_*"],
        ["RQ2 output", "outputs/tables/rq2_incremental_value.csv"],
        ["RQ3 outputs", "outputs/tables/rq3_* and model_selection.json"],
        ["RQ4 outputs", "outputs/tables/rq4_* and outputs/figures/rq4_*"],
        ["Sensitivity run", "outputs/sensitivity/constant_dollars/"],
        ["Official-run provenance", "outputs/official_run_manifest.json"],
    ], columns=["artifact", "repository_path"])
    _table(doc, evidence, [("artifact", "Artifact"), ("repository_path", "Repository Path")])
    _paragraph(doc, "Run sequence: install requirements and the local package; execute the source audit; build the official dataset; run sample-size, RQ1, training, explainability, diagnostics, and report scripts; then run the reproducibility check. The configuration fixes the analysis window, features, holdout, tuning grids, thresholds, and random seed.")

    _heading(doc, "Appendix B")
    _heading(doc, "Decision Rules and Interpretation Safeguards", 2)
    safeguards = [
        "The preferred fitted model must have pooled MASE below 1.0 and MAE below seasonal naive.",
        "A candidate is rejected if it is more than 10% worse than seasonal naive in three or more CMAs.",
        "Macroeconomic variables require at least 5% MAE improvement to establish practical value.",
        "All primary performance claims use the latest 12-month holdout, not training fit.",
        "Building permits are construction intentions rather than completed construction.",
        "Feature importance and regression coefficients are not interpreted causally.",
        "Large projects are reviewed as possible genuine events rather than automatically deleted.",
    ]
    for item in safeguards:
        list_style = next((s for s in doc.styles if s.style_id == "ListBullet"), None)
        p = doc.add_paragraph(style=list_style)
        p.paragraph_format.line_spacing = 2
        r = p.add_run(item)
        _set_font(r)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "QM640 Capstone Final Report Draft"
    doc.core_properties.author = "Debodip Gupta"
    doc.core_properties.comments = "Generated from frozen official analysis outputs; mentor name confirmed from the prior submitted interim report."
    _enforce_times_new_roman(doc)
    doc.save(output)
    return output
