"""Build the final 40-page submission from the supplied template and frozen evidence."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .config import project_root
from .final_report import TITLE, _clear_document, _hyperlink

MODEL_NAMES = {"xgboost": "XGBoost", "random_forest": "Random forest", "elastic_net_macro": "Elastic net + macro",
               "elastic_net_history": "Elastic net history", "seasonal_naive": "Seasonal naive"}
REPO = "https://github.com/dgupta91114/ontario-building-permit-forecasting"

# Author/date prefixes, article titles, italic journal/book titles, and linked identifiers.
REFERENCES = [
    ("Angelopoulos, A. N., & Bates, S. (2023). Conformal prediction: A gentle introduction. ",
     "Foundations and Trends in Machine Learning, 16", "(4), 494–591. ", "https://doi.org/10.1561/2200000101"),
    ("Bank of Canada. (n.d.). ", "Valet API: How-to guide", ". ", "https://www.bankofcanada.ca/valet-api-how-to/"),
    ("Bergmeir, C., Hyndman, R. J., & Koo, B. (2018). A note on the validity of cross-validation for evaluating autoregressive time series prediction. ",
     "Computational Statistics & Data Analysis, 120", ", 70–83. ", "https://doi.org/10.1016/j.csda.2017.11.003"),
    ("Breiman, L. (2001). Random forests. ", "Machine Learning, 45", "(1), 5–32. ", "https://doi.org/10.1023/A:1010933404324"),
    ("Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. In ",
     "Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining",
     " (pp. 785–794). Association for Computing Machinery. ", "https://doi.org/10.1145/2939672.2939785"),
    ("Cohen, J. (1992). A power primer. ", "Psychological Bulletin, 112", "(1), 155–159. ", "https://doi.org/10.1037/0033-2909.112.1.155"),
    ("Diebold, F. X., & Mariano, R. S. (1995). Comparing predictive accuracy. ", "Journal of Business & Economic Statistics, 13", "(3), 253–263. ", "https://doi.org/10.1080/07350015.1995.10524599"),
    ("Goh, B. H. (1996). Residential construction demand forecasting using economic indicators: A comparative study of artificial neural networks and multiple regression. ",
     "Construction Management and Economics, 14", "(1), 25–34. ", "https://doi.org/10.1080/01446199600000004"),
    ("Hyndman, R. J., & Koehler, A. B. (2006). Another look at measures of forecast accuracy. ", "International Journal of Forecasting, 22", "(4), 679–688. ", "https://doi.org/10.1016/j.ijforecast.2006.03.001"),
    ("Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. In I. Guyon, U. von Luxburg, S. Bengio, H. Wallach, R. Fergus, S. Vishwanathan, & R. Garnett (Eds.), ",
     "Advances in neural information processing systems", " (Vol. 30, pp. 4765–4774). Curran Associates. ",
     "https://proceedings.neurips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html"),
    ("Newey, W. K., & West, K. D. (1987). A simple, positive semi-definite, heteroskedasticity and autocorrelation consistent covariance matrix. ",
     "Econometrica, 55", "(3), 703–708. ", "https://doi.org/10.2307/1913610"),
    ("Shmueli, G. (2010). To explain or to predict? ", "Statistical Science, 25", "(3), 289–310. ", "https://doi.org/10.1214/10-STS330"),
    ("Statistics Canada. (2026a). ", "Building permits (BPER): Detailed information for July 2026", ". ",
     "https://www23.statcan.gc.ca/imdb/p2SV.pl?Function=getSurvey&SDDS=2802"),
    ("Statistics Canada. (2026b). ", "Building permits, by type of structure and type of work", " (Table 34-10-0292-01) [Data set]. ",
     "https://doi.org/10.25318/3410029201-eng"),
    ("Statistics Canada. (2026c). ", "Labour force characteristics, monthly, seasonally adjusted and trend-cycle", " (Table 14-10-0287-01) [Data set]. ",
     "https://doi.org/10.25318/1410028701-eng"),
    ("Tashman, L. J. (2000). Out-of-sample tests of forecasting accuracy: An analysis and review. ", "International Journal of Forecasting, 16", "(4), 437–450. ",
     "https://doi.org/10.1016/S0169-2070(00)00065-0"),
    ("Zou, H., & Hastie, T. (2005). Regularization and variable selection via the elastic net. ", "Journal of the Royal Statistical Society: Series B (Statistical Methodology), 67", "(2), 301–320. ",
     "https://doi.org/10.1111/j.1467-9868.2005.00503.x"),
]


def font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "Times New Roman")


class Builder:
    def __init__(self, root):
        self.root = root
        self.out = root / "reports/final"
        self.tables = root / "outputs/tables"
        self.doc = Document(root / "reports/templates/QM_640_Final_Report_Template.docx")
        _clear_document(self.doc)
        # The supplied template does not flag a default paragraph style.
        # Set it explicitly so references inherit the same double leading as body text.
        self.doc.styles["Normal"]._element.set(qn("w:default"), "1")
        self.catalog = []
        self.page_number = 0
        for name in ["Normal", "Heading1", "Heading2", "Heading3"]:
            style = next(s for s in self.doc.styles if s.style_id == name)
            style.font.name = "Times New Roman"; style.font.size = Pt(12)
            style.font.color.rgb = None
            pf = style.paragraph_format
            pf.line_spacing = Pt(24); pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.first_line_indent = Inches(.5) if name == "Normal" else Inches(0)
            pf.keep_with_next = name != "Normal"
            pf.keep_together = False
        for section in self.doc.sections:
            section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
            section.header_distance = Inches(.5)
            header = section.header.paragraphs[0]
            header.clear(); header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header.paragraph_format.first_line_indent = Inches(0)
            header.paragraph_format.line_spacing = 1
            field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
            run = OxmlElement("w:r"); props=OxmlElement("w:rPr")
            rfonts=OxmlElement("w:rFonts"); rfonts.set(qn("w:ascii"),"Times New Roman"); rfonts.set(qn("w:hAnsi"),"Times New Roman")
            size=OxmlElement("w:sz"); size.set(qn("w:val"),"24"); props.extend([rfonts,size]); run.append(props)
            t=OxmlElement("w:t"); t.text="1"; run.append(t); field.append(run); header._p.append(field)
            section.different_first_page_header_footer = False
        self.doc.core_properties.title = TITLE
        self.doc.core_properties.author = "Debodip Gupta"
        self.doc.core_properties.subject = "QM640 Data Analytics Capstone — Final Report"
        self.doc.core_properties.comments = "Final submission; Summer 2026; frozen official empirical evidence."

    def paragraph(self, text, indent=True, size=12, spacing=2):
        p=self.doc.add_paragraph()
        p.paragraph_format.first_line_indent=Inches(.5 if indent else 0)
        p.paragraph_format.line_spacing=Pt(24) if spacing==2 else spacing
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.widow_control=True
        font(p.add_run(text),size)
        return p

    def heading(self, text, level=1):
        style=next(s for s in self.doc.styles if s.style_id == f"Heading{level}")
        p=self.doc.add_paragraph(style=style)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Inches(0)
        font(p.add_run(text),bold=True)
        return p

    def caption(self, label, title):
        p=self.paragraph(label,indent=False,spacing=1.2)
        p.runs[0].bold=True; p.paragraph_format.keep_with_next=True
        p=self.paragraph(title,indent=False,spacing=1.2)
        p.runs[0].italic=True; p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(5)
        self.catalog.append({"label":label,"title":title,"planned_page":self.page_number})

    def note(self,text):
        p=self.doc.add_paragraph(); p.paragraph_format.first_line_indent=Inches(0)
        p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
        font(p.add_run("Note. "),size=10.5,italic=True); font(p.add_run(text),size=10.5)

    def table(self, label, title, headings, rows, note, widths=None, size=10.5):
        self.caption(label,title)
        table=self.doc.add_table(rows=1,cols=len(headings)); table.autofit=False
        width=6.268
        widths=widths or [width/len(headings)]*len(headings)
        for col,w in zip(table.columns,widths): col.width=Inches(w)
        for cell,txt,w in zip(table.rows[0].cells,headings,widths): cell.text=txt; cell.width=Inches(w)
        for row in rows:
            for cell,txt,w in zip(table.add_row().cells,row,widths): cell.text=str(txt); cell.width=Inches(w)
        props=table._tbl.tblPr
        borders=OxmlElement("w:tblBorders")
        for edge in ["top","left","bottom","right","insideH","insideV"]:
            e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single" if edge in ("top","bottom") else "nil")
            e.set(qn("w:sz"),"6"); borders.append(e)
        props.append(borders)
        for i,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i==0:
                repeat=OxmlElement("w:tblHeader"); trpr.append(repeat)
            for cell in row.cells:
                margins=OxmlElement("w:tcMar")
                for edge,val in [("top",45),("bottom",45),("left",55),("right",55)]:
                    e=OxmlElement("w:"+edge); e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa"); margins.append(e)
                cell._tc.get_or_add_tcPr().append(margins)
                if i==0:
                    cb=OxmlElement("w:tcBorders"); e=OxmlElement("w:bottom"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"6"); cb.append(e); cell._tc.get_or_add_tcPr().append(cb)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.line_spacing=1.08
                    p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                    for r in p.runs: font(r,size=size,bold=i==0)
        self.note(note)

    def frame(self,filename): return pd.read_csv(self.tables/filename)

    def data_table(self,key):
        money=lambda x: f"{float(x)/1000:,.3f}"
        dec=lambda x: f"{float(x):.3f}"
        cma=lambda x:x.replace("Kitchener-Cambridge-Waterloo","KCW")
        if key.startswith("literature"):
            cols=["Author (year)","Domain / context","Dataset / setting","Method(s)","Key findings","Relevance to RQ / decision"]
            rows1=[
                ["Goh (1996)","Residential construction demand","Singapore construction and economic indicators","Neural networks; multiple regression","Establishes a direct linear/nonlinear demand comparison; no numerical result imported from unavailable full text.","RQ2–RQ3: test economic inputs and nonlinear alternatives locally."],
                ["Tashman (2000)","Forecast evaluation design","Methodological review and forecasting software","Comparison of origins, windows, and test schemes","Evaluation reliability depends on explicit splitting and multiple forecast origins.","RQ3: use time-ordered folds and specify refitting and test periods."],
                ["Bergmeir et al. (2018)","Autoregressive validation","Theory, simulations, and an empirical example","K-fold and out-of-sample comparisons","Ordinary CV can be valid under particular error conditions; validity is conditional.","RQ3: do not assume those conditions; group complete months."],
                ["Hyndman & Koehler (2006)","Forecast error measurement","Accuracy measures and illustrative series","Comparison of metrics; MASE","Common measures can behave poorly; seasonal scaling supports comparison across series.","RQ3–RQ4: report MASE with dollar errors and region-specific scales."],
                ["Diebold & Mariano (1995)","Predictive accuracy testing","Paired forecasts; statistical theory and examples","Loss-differential tests","Forecast comparison can accommodate general losses and dependent errors when assumptions hold.","RQ2–RQ3: match targets; disclose limitations of pooled tests."],
                ["Cohen (1992)","Statistical power planning","Standard behavioral/statistical test designs","Effect-size and power relationships","Power depends jointly on effect size, sample size, and significance level.","RQ1–RQ4: justify planned counts and disclose actual shortfalls."],
            ]
            rows2=[
                ["Zou & Hastie (2005)","Regularized regression","Simulations and empirical examples","Elastic-net penalty","Combined penalties support shrinkage and grouping of correlated predictors.","RQ2–RQ3: handle correlated lag and rolling predictors with a regularized baseline."],
                ["Breiman (2001)","Ensemble learning","Classification/regression examples and theory","Randomized tree averaging","Tree strength and inter-tree dependence affect ensemble error.","RQ3: evaluate random forest as a nonlinear candidate."],
                ["Chen & Guestrin (2016)","Boosted-tree systems","Machine-learning benchmarks and scalability studies","Regularized gradient boosting","Efficient tree-boosting implementation with sparsity and computational improvements.","RQ3: fit shallow XGBoost with constrained grids and chronological scoring."],
                ["Lundberg & Lee (2017)","Model interpretability","Theory and empirical model explanations","Additive feature attributions (SHAP)","Unifies attribution methods through an additive framework and desirable properties.","RQ4: inspect log-scale contributions without claiming causality."],
                ["Shmueli (2010)","Statistical modeling goals","Conceptual analysis with examples","Distinguishes explanation, prediction, and description","A model's explanatory fit does not establish predictive utility.","All RQs: separate causal language, descriptive regression, and holdout performance."],
                ["Angelopoulos & Bates (2023)","Predictive uncertainty","Tutorial theory and worked applications","Conformal prediction","Calibration can create useful uncertainty sets under stated conditions.","RQ4: report coverage and width; disclose dependence and calibration reuse."],
            ]
            first=key.endswith("1")
            self.table("Table 1" if first else "Table 2","Literature Relevance Matrix: Domain and Evaluation" if first else "Literature Relevance Matrix: Models, Interpretation, and Uncertainty",cols,rows1 if first else rows2,
                       "Original paraphrases of the cited sources. Complete APA references appear in the bibliography. A source's original setting is distinguished from this study's application.",
                       widths=[.78,.86,.91,.91,1.40,1.408],size=9.5)
        elif key=="sources":
            self.table("Table 3","Official Data Sources and Frozen Analytical Scope",["Source","Retained measure","Access / frequency"],[
                ["Statistics Canada 34-10-0292-01","Residential permit value; 5 CMAs","Full-table CSV; monthly"],
                ["Statistics Canada 14-10-0287-01","Ontario unemployment rate","Full-table CSV; monthly"],
                ["Bank of Canada V39079","Target for the overnight rate","Valet JSON; daily → month end"]],
                "All selected monthly series cover January 2018–May 2026. Retrieval: August 16, 2026 (UTC). Sources: Bank of Canada (n.d.) and Statistics Canada (2026b, 2026c).",[2.05,2.2,2.018])
        elif key=="sample":
            self.table("Table 4","Planned Sample Sizes and the Samples Actually Evaluated",["RQ / planning test","Effect assumption","Minimum n","Available / interpretation"],[
                ["RQ1: Fisher-z correlation","r = .30; two-sided","85","88 per CMA; serially dependent; approximation to descriptive design"],
                ["RQ2: regression F test","f² = .15; 10 predictors","118","380 training rows; final paired ablation uses 60 and differs from planning test"],
                ["RQ3: paired t test","d = .35; two-sided","67","60 holdout pairs; below planned count"],
                ["RQ4: 5-group ANOVA","f = .25","196 total","60 total; 12 per CMA; below planned count"]],
                "α = .05; target power = .80. Minima reproduced with SciPy/statsmodels. Rows sharing time or region are not guaranteed independent. RQ4 equal allocation rounds to 200 total.",[1.52,1.16,.7,2.888])
        elif key=="splits":
            self.table("Table 5","Chronological Training, Validation, Holdout, and Calibration Windows",["Stage","Fit target months","Evaluation months","Rows: fit / eval"],[
                ["CV fold 1","Feb 2019–Jan 2022","Feb 2022–Jan 2023","180 / 60"],
                ["CV fold 2","Feb 2019–Jan 2023","Feb 2023–Jan 2024","240 / 60"],
                ["CV fold 3","Feb 2019–Jan 2024","Feb 2024–Jan 2025","300 / 60"],
                ["Point-model refit","Feb 2019–May 2025","Jun 2025–May 2026","380 / 60"],
                ["Interval proper fit","Feb 2019–May 2024","Calibration: Jun 2024–May 2025","320 / 60"],
                ["Interval evaluation","Same proper fit","Jun 2025–May 2026","320 / 60"]],
                "All partitions preserve complete five-CMA months. Model parameters are fixed within each evaluation window; lagged inputs advance by origin month.",[1.35,1.72,2.05,1.148],size=10)
        elif key=="models":
            self.table("Table 6","Candidate Configurations and Selected Hyperparameters",["Specification / inputs","Searched values","Chosen settings"],[
                ["Seasonal naive / annual history","No tuning","Same target calendar month one year earlier"],
                ["Elastic net history / 11 numeric + CMA","α ∈ {.1, 1}; l1 ratio ∈ {.5, 1}","α = .1; l1 ratio = .5"],
                ["Elastic net + macro / 13 numeric + CMA","Same four configurations","α = .1; l1 ratio = .5"],
                ["Random forest / full inputs","150 trees; depth ∈ {3, unrestricted}; leaf ∈ {2, 4}; features = .7","150 trees; depth 3; leaf 2; features .7"],
                ["XGBoost / full inputs","Trees ∈ {100, 200}; depth ∈ {2, 3}; rate .05; row/column sample .8","100 trees; depth 2; rate .05; row/column sample .8"]],
                "Selected settings minimize earlier three-fold mean MAE. All fitted models use log1p targets, fold-fitted preprocessing, and original-scale evaluation.",[1.66,2.50,2.108],size=10)
        elif key=="training":
            f=pd.read_csv(self.out/"evidence/training_validation_test_metrics.csv")
            rows=[[MODEL_NAMES[r.model],money(r.train_mae),money(r.train_rmse),money(r.cv_mae),money(r.test_mae),money(r.test_rmse)] for r in f.itertuples()]
            self.table("Table 7","Training, Cross-Validation, and Holdout Accuracy",["Model","Train MAE","Train RMSE","CV MAE","Test MAE","Test RMSE"],rows,
                       "Monetary errors: CAD millions. Training n = 380; each CV validation fold n = 60; holdout n = 60. Training errors are resubstitution diagnostics, not independent forecast accuracy.",[1.55,.94,.97,.93,.93,.948])
        elif key=="holdout":
            f=self.frame("rq3_model_metrics.csv")
            rows=[[MODEL_NAMES[r.model],money(r.mae),money(r.rmse),dec(r.mase),f"{r.smape_pct:.2f}",f"{r.wape_pct:.2f}"] for r in f.itertuples()]
            self.table("Table 8","Final Holdout Model Comparison",["Model","MAE","RMSE","MASE","sMAPE %","WAPE %"],rows,
                       "n = 60 for every model, June 2025–May 2026. MAE/RMSE in CAD millions; MASE unitless. Lower is better. XGBoost is preferred under the stated acceptance rule.",[1.55,.96,.96,.82,1.0,.978])
        elif key=="summary":
            f=self.frame("rq1_summary_by_cma.csv")
            rows=[[cma(r.cma),money(r.mean),money(r.std),dec(r.coefficient_of_variation)] for r in f.itertuples()]
            self.table("Table 9","Descriptive Regional Permit Values",["CMA","Mean","SD","Coefficient of variation"],rows,
                       "n = 88 origins per CMA, January 2019–April 2026. Mean and SD: CAD millions. SD = standard deviation; KCW = Kitchener-Cambridge-Waterloo.",[1.7,1.3,1.3,1.968])
        elif key=="temporal":
            f=self.frame("rq1_persistence.csv")
            rows=[[cma(r.cma),dec(r.acf_lag_1),dec(r.acf_lag_12)] for r in f.itertuples()]
            self.table("Table 10","Regional Autocorrelation and Pooled Temporal Tests",["CMA","Lag-1 ACF","Lag-12 ACF"],rows,
                       "ACF = autocorrelation function. Pooled OLS: current permit b = .315, p < .001; time b = 575.298 (CAD thousands/month), p = .007; sine p = .998; cosine p = .708; R² = .923. Nominal HAC tests use three pooled-row lags.",[2.45,1.9,1.918])
        elif key=="rq2":
            r=self.frame("rq2_incremental_value.csv").iloc[0]
            self.table("Table 11","Incremental Value of Macroeconomic Inputs in Elastic Net",["Quantity","Result"],[
                ["History-only MAE",money(r.mae_b)+" million CAD"],
                ["Macro-augmented MAE",money(r.mae_a)+" million CAD"],
                ["Relative MAE improvement",f"{r.improvement_pct_a_vs_b:.3f}%"],
                ["Practical threshold","5%; not met"],
                ["Paired t / Wilcoxon / DM-style p",".280 / .276 / .256"]],
                "Original frozen comparison: 60 paired CMA-month errors. Positive improvement favors macro augmentation. All formal decisions use nominal α = .05.",[3.65,2.618])
        elif key=="regional":
            f=self.frame("rq4_regional_metrics.csv"); x=f[f.model=="xgboost"]; n=f[f.model=="seasonal_naive"].set_index("cma")
            rows=[[cma(r.cma),money(r.mae),dec(r.mase),f"{100*(n.loc[r.cma,'mae']-r.mae)/n.loc[r.cma,'mae']:.1f}%"] for r in x.itertuples()]
            self.table("Table 12","XGBoost Regional Holdout Performance",["CMA","MAE (CAD m)","MASE","MAE reduction vs. naive"],rows,
                       "n = 12 per CMA. KCW = Kitchener-Cambridge-Waterloo. Every CMA improves against the same-target seasonal-naive control.",[1.58,1.35,1.05,2.288])
        elif key=="monthly":
            f=pd.read_csv(self.out/"evidence/supplementary_month_level_comparisons.csv"); f=f[f.model_a=="xgboost"]
            rows=[["XGBoost vs. "+MODEL_NAMES[r.model_b],money(r.mean_mae_reduction),f"{money(r.ci95_lower)} to {money(r.ci95_upper)}",dec(r.p_value),f"{r.positive_months}/12"] for r in f.itertuples()]
            self.table("Table 13","Supplementary Comparison of Monthly Pooled Losses",["Comparison","Mean gain","95% interval","p","Months better"],rows,
                       "Gains/intervals: CAD millions. Twelve monthly mean differences; t tests with 11 df. Exploratory: serial dependence is not removed by monthly aggregation.",[1.9,.86,1.68,.68,1.148],size=10)
        elif key=="coverage":
            f=self.frame("rq4_interval_coverage_by_cma.csv")
            rows=[[cma(r.cma),"12/12",f"{r.coverage:.0%}",money(r.mean_interval_width)] for r in f.itertuples()]
            self.table("Table 14","Empirical Coverage and Width of the Separate Interval Model",["CMA","Covered","Coverage","Mean width (CAD m)"],rows,
                       "Nominal level: 90%. These intervals and their point centers come from the 320-row proper fit with 60 calibration rows, not the 380-row headline refit.",[1.85,1.05,1.15,2.218])
        elif key=="sensitivity":
            f=pd.read_csv(self.root/"outputs/sensitivity/constant_dollars/outputs/tables/rq3_model_metrics.csv")
            rows=[[MODEL_NAMES[r.model],money(r.mae),money(r.rmse),dec(r.mase)] for r in f.itertuples()]
            self.table("Table 15","Constant-Dollar Sensitivity: Holdout Performance",["Model","MAE","RMSE","MASE"],rows,
                       "n = 60 per model; same dates and CMA scope. Monetary errors: millions of constant CAD on the official source basis. Compare models within this table; current/constant monetary errors have different value bases.",[2.35,1.3,1.3,1.318])
        elif key=="example":
            f=self.frame("selected_model_intervals.csv"); r=f[(f.cma=="Kitchener-Cambridge-Waterloo") & (f.target_date=="2025-06-01")].iloc[0]
            self.table("Table 16","Retrospective Scorecard Example: KCW, June 2025",["Field","Recorded value"],[
                ["Interval-model point",money(r.prediction)+" million CAD"],
                ["Lower / upper 90% bounds",money(r.lower_90)+" / "+money(r.upper_90)+" million CAD"],
                ["Subsequently observed value",money(r.actual)+" million CAD"],
                ["Example status","Historical illustration; proposed user interface"]],
                "All three forecast fields are from the same interval-model record. KCW = Kitchener-Cambridge-Waterloo. No live forecast or measured user trial is represented.",[2.6,3.668])
        elif key=="dictionary":
            rows=[
                ["ref_date; target_date","Monthly dates","Key","Origin reference month; target is one month later."],
                ["cma","Category","Predictor/key","One of the five specified CMAs; one-hot encoded."],
                ["dguid; source_status","Text / flag","Audit","Official geography identifier and published quality flag."],
                ["permit_value_t","CAD thousands","Predictor","Observed origin-month residential permit value."],
                ["permit_value_t_plus_1","CAD thousands","Target","Next-month value within the same CMA."],
                ["log_target","Log units","Target transform","log(1 + target); excluded from input features."],
                ["lag_1; lag_3; lag_6; lag_12","CAD thousands","Predictors","Permit values at t−1, t−3, t−6, t−12."],
                ["rolling_mean_3","CAD thousands","Predictor","Mean of values at t, t−1, t−2."],
                ["rolling_sd_3","CAD thousands","Predictor","Sample SD across t, t−1, t−2; ddof = 1."],
                ["seasonal_naive_pred","CAD thousands","Benchmark","y(t−11), aligned to target t+1 one year earlier."],
                ["month_sin; month_cos","−1 to 1","Predictors","Sine/cosine of target calendar month with period 12."],
                ["time_index","Months","Predictor","Elapsed origin months from January 2018."],
                ["pandemic_indicator","0 / 1","Predictor","1 for targets March 2020–June 2021."],
                ["overnight_rate","Percent","Source/audit","Last V39079 observation in origin month."],
                ["ontario_unemployment_rate","Percent","Source/audit","Ontario adjusted unemployment, age 15+, total gender."],
                ["overnight_rate_lag_1; unemployment_lag_1","Percent","Macro predictors","Previous-month source rates; lagged before CMA replication."],
            ]
            self.table("Table A1","Modeling Dataset Fields and Analytical Roles",["Field(s)","Unit / type","Role","Definition"],rows,
                       "Derived from the actual feature code and the 23-column processed dataset. Macro predictors are included only in the augmented elastic net and tree models.",[1.82,.88,.91,2.658],size=9.5)
        elif key=="evidence":
            rows=[
                ["Study scope and model settings","config/config.yaml"],
                ["Official source hashes and retrieval dates","outputs/official_run_manifest.json"],
                ["Complete analytic data; construction audit","data/processed/modeling_dataset.csv; feature_audit.json"],
                ["Source filtering evidence","reports/final/evidence/cleaning_audit.json"],
                ["RQ1 / RQ2 evidence","outputs/tables/rq1_*; rq2_incremental_value.csv"],
                ["RQ3 predictions, metrics, selection","outputs/tables/holdout_predictions.csv; rq3_*; model_selection.json"],
                ["RQ4 reliability, attribution, intervals","outputs/tables/rq4_*; selected_model_intervals.csv"],
                ["Alternative value basis","outputs/sensitivity/constant_dollars/"],
                ["Final-report supplementary diagnostics","reports/final/evidence/"],
                ["Manuscript; final validation","reports/final/manuscript.md; quality_audit.json"],
            ]
            self.table("Table B1","Repository Evidence and Reproduction Guide",["Evidence","Repository location"],rows,
                       "All paths are relative to the GitHub repository linked after the title page. Raw archives and fitted binaries are local regeneration inputs; compact data, outputs, and report diagnostics are versioned.",[2.33,3.938],size=10)
        else: raise KeyError(key)

    def figure(self,key):
        mapping={
            "architecture":("Figure 1","End-to-End Workflow of the Proposed System",self.out/"evidence/architecture.png",6.1,"Original workflow diagram. Solid = implemented analysis; dashed = proposed pilot; arrows = artifact flow. The two fitting branches produce different point predictions."),
            "comparison":("Figure 2","Holdout Error by Model",self.out/"evidence/model_comparison.png",6.2,"Original visualization of Table 8. Dollar and scaled errors provide complementary rankings; all models share the same target observations."),
            "history":("Figure 3","Regional Permit Histories Used as Model Inputs",self.out/"evidence/regional_history.png",6.2,"Original plot from the processed official dataset. KCW = Kitchener-Cambridge-Waterloo. Seasonally adjusted current values, CAD millions; separate y-axis scales."),
            "forecasts":("Figure 4","XGBoost Point Forecasts and Observed Holdout Values",self.out/"evidence/holdout_forecasts.png",6.2,"Original plot of frozen point-model predictions, CAD millions. Target months: June 2025–May 2026. KCW = Kitchener-Cambridge-Waterloo. These are not interval-model centers."),
            "permutation":("Figure 5","Leading Permutation-Importance Features",self.out/"evidence/permutation.png",6.15,"Original visualization of saved holdout permutation estimates. Higher values indicate more damage to accuracy after an input is shuffled."),
            "shap":("Figure 6","Leading Mean Absolute SHAP Attributions",self.out/"evidence/shap_importance.png",6.15,"Original plot of archived SHAP summaries. Units are log1p-target contributions; mean absolute attribution conveys magnitude, not direction or causality."),
        }
        label,title,path,width,note=mapping[key]
        self.caption(label,title)
        p=self.doc.add_paragraph(); p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.line_spacing=1
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
        p.paragraph_format.keep_with_next=True
        self.note(note)

    def references(self,part):
        # Manual page breaks keep individual APA entries intact.
        ranges={1:(0,6),2:(6,12),3:(12,17)}
        start,end=ranges[part]
        for lead,italic,tail,url in REFERENCES[start:end]:
            p=self.doc.add_paragraph(); pf=p.paragraph_format
            pf.left_indent=Inches(.5); pf.first_line_indent=Inches(-.5); pf.keep_together=True
            font(p.add_run(lead)); font(p.add_run(italic),italic=True); font(p.add_run(tail)); _hyperlink(p,url,url)

    def build(self):
        manuscript=(self.out/"manuscript.md").read_text(encoding="utf-8")
        pages=manuscript.split("<!-- PAGE -->")
        assert len(pages)==40,len(pages)
        page_map=[]
        for number,page in enumerate(pages,1):
            self.page_number=number
            first_paragraph_index=len(self.doc.paragraphs)
            page_map.append({"planned_page":number,"heading":next((x.lstrip("# ") for x in page.splitlines() if x.startswith("#")),"Title page" if number==1 else f"Evidence / references {number}"),"manuscript_words":len(page.split())})
            for block in re.split(r"\n\s*\n",page.strip()):
                block=block.strip()
                if not block: continue
                if block=="[[TITLE_PAGE]]":
                    p=self.paragraph("",indent=False); p.paragraph_format.space_after=Pt(60)
                    for txt,bold in [("Data Analytics Capstone",True),(TITLE,True),("Final Report",True),("Debodip Gupta",False),("Walsh College",False),("QM640: Data Analytics Capstone",False),("Dr. Sanhita Karmakar",False),("Summer 2026 Term",False),("September 15, 2026",False)]:
                        p=self.paragraph(txt,indent=False); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=bold
                    continue
                if block=="[[REPOSITORY]]":
                    p=self.paragraph("GitHub repository: ",indent=False,size=11,spacing=1.2)
                    _hyperlink(p,REPO.replace("https://",""),REPO); p.paragraph_format.space_after=Pt(6)
                    continue
                match=re.fullmatch(r"\[\[(TABLE|FIGURE|REFERENCES):([^]]+)\]\]",block)
                if match:
                    kind,key=match.groups()
                    {"TABLE":self.data_table,"FIGURE":self.figure,"REFERENCES":lambda x:self.references(int(x))}[kind](key)
                elif block.startswith("#"):
                    marks,text=block.split(" ",1); self.heading(text,len(marks))
                else:
                    self.paragraph(" ".join(block.splitlines()),indent=not (number==2 or block.startswith("Note.")))
            if number>1:
                self.doc.paragraphs[first_paragraph_index].paragraph_format.page_break_before=True
        path=self.out/"QM640_Capstone_Final_Report_Debodip_Gupta.docx"
        self.doc.save(path)
        (self.out/"page_map.json").write_text(json.dumps(page_map,indent=2),encoding="utf-8")
        (self.out/"figure_table_catalog.json").write_text(json.dumps(self.catalog,indent=2),encoding="utf-8")
        (self.out/"references.json").write_text(json.dumps(REFERENCES,indent=2),encoding="utf-8")
        return path


def generate_final_report(root: Path | None=None):
    return Builder(root or project_root()).build()
